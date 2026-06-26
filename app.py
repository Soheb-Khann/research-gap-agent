import streamlit as st
import tempfile, os, sys
from pathlib import Path
from src.ingestion.embedder import reset_collection
import re
import io


st.set_page_config(page_title="Research Gap Agent", page_icon="🔬")
st.title("🔬 Research Gap Agent")

uploaded = st.file_uploader("Upload PDFs", type=["pdf"], accept_multiple_files=True)

if st.button("▶ Run Pipeline", type="primary", disabled=not uploaded):

    # Save uploads to temp dir
    tmp_dir = tempfile.mkdtemp()
    pdf_paths = []
    for uf in uploaded:
        dest = os.path.join(tmp_dir, uf.name)
        with open(dest, "wb") as f:
            f.write(uf.getbuffer())
        pdf_paths.append(dest)

    # Add project root so src.* imports work
    root = Path(__file__).parent
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))

    from src.graph.orchestrator import build_graph

    with st.status("Running pipeline...", expanded=True) as status:
        st.write("🧹 Clearing previous session data...")
        reset_collection() 
        st.write("⏳ Step 1/5 — Ingesting and embedding PDFs...")
        st.write("⏳ Step 2/5 — Summarising each paper (this takes the longest)...")
        st.write("⏳ Step 3/5 — Running cross-paper analysis...")
        st.write("⏳ Step 4/5 — Identifying research gaps...")
        st.write("⏳ Step 5/5 — Generating report...")
        graph = build_graph()
        result = graph.invoke({
            "pdf_paths": pdf_paths,
            "chunks": [], "summaries": [], "analysis": {},
            "gaps": [], "final_report": "", "error": ""
        })
        if result.get("error"):
            status.update(label="❌ Pipeline failed", state="error")
            st.error(result["error"])
            st.stop()
        status.update(label="✅ Done!", state="complete")
    st.session_state["result"] = result

if "result" in st.session_state:
    result    = st.session_state["result"]
    gaps      = result.get("gaps", [])
    summaries = result.get("summaries", [])
    report    = result.get("final_report", "")

    # Stats
    st.divider()
    col_a, col_b, col_c = st.columns(3)
    col_a.metric("Papers analysed", len(summaries))
    col_b.metric("Gaps identified", len(gaps))
    col_c.metric("High severity",   sum(1 for g in gaps if g.get("severity") == "high"))

    # Report
    st.divider()
    st.markdown(report)
    st.divider()

    # Downloads
    col1, col2 = st.columns(2)

    with col1:
        st.download_button(
            "⬇ Download Markdown",
            data=report,
            file_name="research_gap_report.md",
            mime="text/markdown",
            use_container_width=True
        )

    with col2:
        try:
            from docx import Document

            doc = Document()
            for line in report.split("\n"):
                line = line.rstrip()
                if   line.startswith("# "):    doc.add_heading(line[2:], 1)
                elif line.startswith("## "):   doc.add_heading(line[3:], 2)
                elif line.startswith("### "):  doc.add_heading(line[4:], 3)
                elif line.startswith("#### "): doc.add_heading(line[5:], 4)
                elif line.startswith("- "):    doc.add_paragraph(line[2:], style="List Bullet")
                elif line == "---":            pass
                elif line:
                    clean = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
                    clean = re.sub(r"_(.*?)_",       r"\1", clean)
                    clean = clean.lstrip("> ")
                    if clean.strip():
                        doc.add_paragraph(clean)

            buf = io.BytesIO()
            doc.save(buf)

            st.download_button(
                "⬇ Download Word (.docx)",
                data=buf.getvalue(),
                file_name="research_gap_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        except ImportError:
            st.info("Install `python-docx` to enable Word export.")

    # Allow user to clear and start fresh
    if st.button("🔄 Analyse new papers"):
        del st.session_state["result"]
        st.rerun()